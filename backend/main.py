"""
main.py - FastAPI application for the ToolChain AI Dashboard (v2 - Stripe billing).

Routes:
  GET  /                       → serve dashboard HTML
  POST /api/auth/signup        → register new user
  POST /api/auth/login         → authenticate user
  GET  /api/auth/me            → current user profile
  GET  /api/auth/username/check → username availability
  GET  /api/auth/account/devices → connected devices
  GET  /api/auth/security → security settings summary
  POST /api/auth/security/2fa/setup → begin 2FA setup
  PUT  /api/auth/security/lock-pin → set website lock PIN
  POST /api/auth/account/devices → register current device
  GET  /api/mcps               → list all registered MCPs
  POST /api/mcps               → register a new MCP
  GET  /api/mcps/{id}          → get single MCP
  DELETE /api/mcps/{id}        → delete an MCP
  POST /api/mcps/{id}/connect  → mark MCP as connected
  POST /api/mcps/{id}/disconnect → disconnect MCP
  POST /api/mcps/{id}/probe    → probe MCP for tools
  GET  /api/mcps/{id}/auth/url → get OAuth URL from MCP server (generic proxy)
  GET  /api/mcps/{id}/auth/status → check MCP auth status (generic proxy)
  POST /api/mcps/{id}/auth/revoke → revoke MCP auth (generic proxy)
  POST /api/sessions           → create chat session
  GET  /api/sessions/{id}/messages → get session history
  POST /api/chat/stream        → SSE streaming chat endpoint
"""
import logging
from contextlib import asynccontextmanager
from typing import Optional

import httpx
from fastapi import FastAPI, HTTPException, Depends, File, UploadFile
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession

from db_config import init_db, get_db, AsyncSessionLocal
from database import (
    register_mcp, list_mcps, get_mcp,
    set_mcp_connection, delete_mcp, get_connected_mcps,
    create_session, list_user_sessions, update_session_title,
    delete_session, save_message, get_session_messages,
    get_all_user_sessions_with_messages, delete_all_user_sessions,
)
from mcp_manager import get_mcp_tools, probe_mcp, execute_tool
from agent import stream_agent_response
from control_room import classify_turn_intent, apply_turn_routing
from agent import build_llm
from chat_attachments import (
    verify_session_owner,
    save_attachment,
    list_attachments,
    delete_attachment,
    get_attachment_file,
    build_agent_attachment_context,
    format_stored_user_message,
)
from auth_routes import router as auth_router
from stripe_routes import router as stripe_router, get_user_limits
from auth_utils import get_current_user, send_contact_email
from plan_guard import (
    require_active_subscription,
    check_daily_message_limit,
    check_session_limit,
    check_mcp_register_limit,
    check_mcp_connect_limit,
    check_tool_execution_access,
    get_usage_stats,
    get_limits,
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
logger = logging.getLogger(__name__)


# ─── Lifespan ────────────────────────────────────────────────────────────────

@asynccontextmanager
async def lifespan(app: FastAPI):
    import asyncio
    await init_db()
    logger.info("PostgreSQL database initialized via SQLAlchemy.")

    async def _expiry_loop():
        """Background task that checks for expired subscriptions every 6 hours."""
        from plan_guard import check_and_downgrade_expired_users
        while True:
            try:
                await asyncio.sleep(6 * 3600)
                async with AsyncSessionLocal() as db:
                    count = await check_and_downgrade_expired_users(db)
                    if count:
                        logger.info("Expiry check: downgraded %d user(s)", count)
            except asyncio.CancelledError:
                break
            except Exception:
                logger.exception("Expiry check failed")

    task = asyncio.create_task(_expiry_loop())
    yield
    task.cancel()
    logger.info("Shutting down.")


# ─── App ─────────────────────────────────────────────────────────────────────
app = FastAPI(
    title="ToolChain AI Dashboard",
    description="Register MCP servers and chat with an AI agent that uses them as tools.",
    version="1.0.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/assets", StaticFiles(directory="../frontend/dist/assets"), name="assets")
app.include_router(auth_router)
app.include_router(stripe_router)


# ─── Pydantic Models ──────────────────────────────────────────────────────────

class MCPCreate(BaseModel):
    name: str
    url: str
    transport: str = "sse"
    description: Optional[str] = None
    icon: Optional[str] = None


class ChatRequest(BaseModel):
    session_id: str
    message: str
    mcp_ids: Optional[list[str]] = None
    attachment_ids: Optional[list[str]] = None


class ContactRequest(BaseModel):
    name: str
    email: str
    subject: str
    category: str
    message: str


class ToolExecuteRequest(BaseModel):
    args: dict = {}


# ─── Routes: Public ──────────────────────────────────────────────────────────

@app.post("/api/contact")
async def api_contact(body: ContactRequest):
    """Public contact form — sends email to admin."""
    if not body.name.strip() or not body.email.strip() or not body.subject.strip() or not body.message.strip():
        raise HTTPException(status_code=400, detail="All fields are required.")
    if body.category not in ("complaint", "feedback", "query"):
        raise HTTPException(status_code=400, detail="Invalid category.")
    try:
        send_contact_email(
            name=body.name.strip(),
            email=body.email.strip(),
            subject=body.subject.strip(),
            category=body.category,
            message=body.message.strip(),
        )
    except RuntimeError as exc:
        raise HTTPException(status_code=500, detail=str(exc))
    return {"message": "Your message has been sent. We'll get back to you soon!"}


# ─── Routes: Dashboard ───────────────────────────────────────────────────────

@app.get("/", include_in_schema=False)
@app.get("/dashboard", include_in_schema=False)
@app.get("/mcp-servers", include_in_schema=False)
@app.get("/tool-execution", include_in_schema=False)
@app.get("/chat", include_in_schema=False)
@app.get("/pricing", include_in_schema=False)
@app.get("/settings", include_in_schema=False)
@app.get("/login", include_in_schema=False)
@app.get("/signup", include_in_schema=False)
@app.get("/forgot-password", include_in_schema=False)
@app.get("/reset-password", include_in_schema=False)
async def serve_frontend():
    return FileResponse("../frontend/dist/index.html")


# ─── Helpers ──────────────────────────────────────────────────────────────────

def _mcp_base_url(mcp_url: str) -> str:
    """Extract base URL from MCP SSE endpoint (strip /sse suffix)."""
    url = mcp_url.rstrip("/")
    if url.endswith("/sse"):
        url = url[:-4]
    return url


# ─── Routes: MCP Registry ────────────────────────────────────────────────────

@app.get("/api/mcps")
async def api_list_mcps(
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    return await list_mcps(db, user["id"])


@app.post("/api/mcps", status_code=201)
async def api_register_mcp(
    body: MCPCreate,
    user: dict = Depends(check_mcp_register_limit),
    db: AsyncSession = Depends(get_db),
):
    existing = await list_mcps(db, user["id"])
    if any(m["name"] == body.name for m in existing):
        raise HTTPException(status_code=409, detail=f"MCP with name '{body.name}' already exists.")

    url = body.url.rstrip("/")
    if body.transport == "sse" and not url.endswith("/sse"):
        url += "/sse"

    probe_result = await probe_mcp(url, body.transport)
    if not probe_result["ok"]:
        raise HTTPException(
            status_code=422,
            detail="Server unreachable. Please check the URL and ensure the server is running.",
        )

    mcp = await register_mcp(db, user["id"], body.name, url, body.transport, body.description, body.icon)
    return mcp


@app.get("/api/mcps/{mcp_id}")
async def api_get_mcp(
    mcp_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    mcp = await get_mcp(db, mcp_id, user["id"])
    if not mcp:
        raise HTTPException(status_code=404, detail="MCP not found.")
    return mcp


@app.delete("/api/mcps/{mcp_id}", status_code=204)
async def api_delete_mcp(
    mcp_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    mcp = await get_mcp(db, mcp_id, user["id"])
    if not mcp:
        raise HTTPException(status_code=404, detail="MCP not found.")

    base = _mcp_base_url(mcp["url"])
    try:
        async with httpx.AsyncClient(timeout=8) as client:
            await client.post(f"{base}/auth/revoke", params={"user_id": user["id"]})
    except Exception:
        logger.debug("Auth revoke skipped for %s (server unreachable)", mcp.get("name", mcp_id))

    await delete_mcp(db, mcp_id, user["id"])


@app.post("/api/mcps/{mcp_id}/connect")
async def api_connect_mcp(
    mcp_id: str,
    skip_auth: bool = False,
    user: dict = Depends(check_mcp_connect_limit),
    db: AsyncSession = Depends(get_db),
):
    mcp = await get_mcp(db, mcp_id, user["id"])
    if not mcp:
        raise HTTPException(status_code=404, detail="MCP not found.")

    if not skip_auth:
        base = _mcp_base_url(mcp["url"])
        try:
            async with httpx.AsyncClient(timeout=8) as client:
                resp = await client.get(f"{base}/auth/status", params={"user_id": user["id"]})
                if resp.status_code == 200:
                    data = resp.json()
                    if not data.get("authenticated"):
                        url_resp = await client.get(f"{base}/auth/url", params={"user_id": user["id"]})
                        if url_resp.status_code == 200:
                            auth_data = url_resp.json()
                            return {
                                "id": mcp_id,
                                "connected": False,
                                "needs_auth": True,
                                "auth_url": auth_data.get("auth_url", ""),
                            }
        except (httpx.RequestError, httpx.HTTPStatusError):
            pass
        except Exception as e:
            logging.warning("Auth check for MCP %s failed: %s", mcp_id, e)

    await set_mcp_connection(db, mcp_id, user["id"], True)
    return {"id": mcp_id, "connected": True}


@app.post("/api/mcps/{mcp_id}/disconnect")
async def api_disconnect_mcp(
    mcp_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    mcp = await get_mcp(db, mcp_id, user["id"])
    if not mcp:
        raise HTTPException(status_code=404, detail="MCP not found.")

    base = _mcp_base_url(mcp["url"])
    try:
        async with httpx.AsyncClient(timeout=8) as client:
            await client.post(f"{base}/auth/revoke", params={"user_id": user["id"]})
    except Exception:
        logger.debug("Auth revoke skipped for %s (server unreachable)", mcp.get("name", mcp_id))

    await set_mcp_connection(db, mcp_id, user["id"], False)
    return {"id": mcp_id, "connected": False}


@app.post("/api/mcps/{mcp_id}/toggle")
async def api_toggle_mcp(
    mcp_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Toggle MCP connection on/off. Used by chat panel."""
    mcp = await get_mcp(db, mcp_id, user["id"])
    if not mcp:
        raise HTTPException(status_code=404, detail="MCP not found.")

    new_state = not mcp.get("connected", False)
    await set_mcp_connection(db, mcp_id, user["id"], new_state)
    return {"id": mcp_id, "connected": new_state}


@app.post("/api/mcps/{mcp_id}/probe")
async def api_probe_mcp(
    mcp_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Probe MCP endpoint to check reachability and list available tools."""
    mcp = await get_mcp(db, mcp_id, user["id"])
    if not mcp:
        raise HTTPException(status_code=404, detail="MCP not found.")
    result = await probe_mcp(mcp["url"], mcp["transport"])
    return result


@app.post("/api/mcps/{mcp_id}/tools/{tool_name}/execute")
async def api_execute_tool(
    mcp_id: str,
    tool_name: str,
    body: ToolExecuteRequest,
    user: dict = Depends(check_tool_execution_access),
    db: AsyncSession = Depends(get_db),
):
    """Directly invoke a single tool on the MCP server (bypasses the agent). Pro+ only."""
    mcp = await get_mcp(db, mcp_id, user["id"])
    if not mcp:
        raise HTTPException(status_code=404, detail="MCP not found.")
    if not mcp.get("connected"):
        raise HTTPException(status_code=400, detail="MCP is not connected.")

    result = await execute_tool(mcp, tool_name, body.args, user_id=str(user["id"]))
    if not result["ok"]:
        raise HTTPException(status_code=422, detail=result["error"])
    return result


# ─── Routes: Generic MCP Auth Proxy ──────────────────────────────────────────
# These forward auth requests to the MCP server's own HTTP auth routes.
# The backend has ZERO knowledge of what kind of auth the MCP uses.

@app.get("/api/mcps/{mcp_id}/auth/url")
async def api_mcp_auth_url(
    mcp_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Get OAuth URL from MCP server. Passes user_id so the MCP can track per-user tokens."""
    mcp = await get_mcp(db, mcp_id, user["id"])
    if not mcp:
        raise HTTPException(status_code=404, detail="MCP not found.")

    base = _mcp_base_url(mcp["url"])
    try:
        async with httpx.AsyncClient(timeout=10) as client:
            resp = await client.get(f"{base}/auth/url", params={"user_id": user["id"]})
            if resp.status_code != 200:
                raise HTTPException(status_code=502, detail="MCP server auth endpoint unreachable.")
            return resp.json()
    except httpx.RequestError as e:
        raise HTTPException(status_code=502, detail=f"Cannot reach MCP auth: {e}")


@app.get("/api/mcps/{mcp_id}/auth/status")
async def api_mcp_auth_status(
    mcp_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Check if the current user is authenticated on this MCP server."""
    mcp = await get_mcp(db, mcp_id, user["id"])
    if not mcp:
        raise HTTPException(status_code=404, detail="MCP not found.")

    base = _mcp_base_url(mcp["url"])
    try:
        async with httpx.AsyncClient(timeout=10) as client:
            resp = await client.get(f"{base}/auth/status", params={"user_id": user["id"]})
            if resp.status_code != 200:
                return {"authenticated": False, "email": None}
            return resp.json()
    except httpx.RequestError:
        return {"authenticated": False, "email": None}


@app.post("/api/mcps/{mcp_id}/auth/revoke")
async def api_mcp_auth_revoke(
    mcp_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Revoke auth for the current user on this MCP server."""
    mcp = await get_mcp(db, mcp_id, user["id"])
    if not mcp:
        raise HTTPException(status_code=404, detail="MCP not found.")

    base = _mcp_base_url(mcp["url"])
    try:
        async with httpx.AsyncClient(timeout=10) as client:
            resp = await client.post(f"{base}/auth/revoke", params={"user_id": user["id"]})
            return resp.json()
    except httpx.RequestError:
        return {"revoked": False}


# ─── Routes: Dashboard Stats ─────────────────────────────────────────────────

@app.get("/api/stats/weekly")
async def api_weekly_stats(
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Get daily chat/message counts for the last 7 days (real-time stats)."""
    from sqlalchemy import func, cast, Date, select
    from db_models import Message, ChatSession
    from datetime import datetime, timedelta

    today = datetime.utcnow().date()
    week_ago = today - timedelta(days=6)

    msg_result = await db.execute(
        select(
            cast(Message.created_at, Date).label("day"),
            func.count(Message.id).label("count"),
        )
        .join(ChatSession, Message.session_id == ChatSession.id)
        .where(
            ChatSession.user_id == user["id"],
            Message.role == "user",
            cast(Message.created_at, Date) >= week_ago,
        )
        .group_by(cast(Message.created_at, Date))
        .order_by(cast(Message.created_at, Date))
    )
    msg_rows = msg_result.all()

    sess_result = await db.execute(
        select(
            cast(ChatSession.created_at, Date).label("day"),
            func.count(ChatSession.id).label("count"),
        )
        .where(
            ChatSession.user_id == user["id"],
            cast(ChatSession.created_at, Date) >= week_ago,
        )
        .group_by(cast(ChatSession.created_at, Date))
        .order_by(cast(ChatSession.created_at, Date))
    )
    sess_rows = sess_result.all()

    msg_map = {str(row.day): row.count for row in msg_rows}
    sess_map = {str(row.day): row.count for row in sess_rows}

    days_data = []
    for i in range(7):
        d = week_ago + timedelta(days=i)
        day_str = str(d)
        day_label = d.strftime("%a")
        days_data.append({
            "day": day_label,
            "date": day_str,
            "messages": msg_map.get(day_str, 0),
            "sessions": sess_map.get(day_str, 0),
        })

    total_messages = sum(d["messages"] for d in days_data)
    total_sessions = sum(d["sessions"] for d in days_data)

    return {
        "days": days_data,
        "total_messages": total_messages,
        "total_sessions": total_sessions,
    }


# ─── Routes: Sessions ────────────────────────────────────────────────────────

@app.get("/api/sessions")
async def api_list_sessions(
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List all chat sessions for the current user."""
    return await list_user_sessions(db, user["id"])


@app.post("/api/sessions", status_code=201)
async def api_create_session(
    user: dict = Depends(check_session_limit),
    db: AsyncSession = Depends(get_db),
):
    session = await create_session(db, user["id"])
    return session


@app.delete("/api/sessions", status_code=200)
async def api_delete_all_sessions(
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Delete all chat sessions and their messages for the current user."""
    count = await delete_all_user_sessions(db, user["id"])
    return {"deleted": count}


@app.get("/api/sessions/export")
async def api_export_sessions(
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Export all chat sessions as a Word .docx file."""
    import io
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from datetime import datetime

    sessions = await get_all_user_sessions_with_messages(db, user["id"])

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title_para = doc.add_heading("Chat Export \u2014 ToolChain AI", level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    username = user.get("username") or user.get("email") or "User"
    doc.add_paragraph(f"Exported by: {username}")
    doc.add_paragraph(f"Date: {datetime.utcnow().strftime('%B %d, %Y at %H:%M UTC')}")
    doc.add_paragraph(f"Total sessions: {len(sessions)}")

    if not sessions:
        doc.add_paragraph("\nNo chat sessions found.")
    else:
        for idx, session in enumerate(sessions, 1):
            created = session.get("created_at", "")
            if created:
                try:
                    dt = datetime.fromisoformat(str(created).replace("Z", "+00:00"))
                    date_str = dt.strftime("%B %d, %Y")
                except Exception:
                    date_str = str(created)[:10]
            else:
                date_str = "Unknown date"

            title_text = session.get("title") or f"Session {idx}"
            doc.add_heading(f"{title_text} ({date_str})", level=2)

            messages = session.get("messages", [])
            if not messages:
                doc.add_paragraph("(No messages)")
            else:
                for msg in messages:
                    role = msg.get("role", "unknown")
                    content = msg.get("content", "")
                    label = "You" if role == "user" else "Assistant" if role == "assistant" else role.capitalize()

                    p = doc.add_paragraph()
                    run = p.add_run(f"[{label}]: ")
                    run.bold = True
                    if role == "user":
                        run.font.color.rgb = RGBColor(0x33, 0x33, 0xCC)
                    else:
                        run.font.color.rgb = RGBColor(0x10, 0xA3, 0x7F)
                    p.add_run(content)

            if idx < len(sessions):
                doc.add_paragraph("\u2500" * 50)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=toolchain_chats_export.docx"},
    )


@app.patch("/api/sessions/{session_id}")
async def api_update_session(
    session_id: str,
    body: dict,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Update session title."""
    title = body.get("title")
    if title:
        await update_session_title(db, session_id, title)
    return {"id": session_id, "title": title}


@app.delete("/api/sessions/{session_id}", status_code=204)
async def api_delete_session(
    session_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Delete a chat session and its messages."""
    deleted = await delete_session(db, session_id, user["id"])
    if not deleted:
        raise HTTPException(status_code=404, detail="Session not found.")


@app.get("/api/sessions/{session_id}/messages")
async def api_get_messages(
    session_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    messages = await get_session_messages(db, session_id)
    return messages


# ─── Routes: Chat attachments ────────────────────────────────────────────────

@app.get("/api/sessions/{session_id}/attachments")
async def api_list_attachments(
    session_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await verify_session_owner(db, session_id, user["id"])
    return list_attachments(user["id"], session_id)


@app.post("/api/sessions/{session_id}/attachments")
async def api_upload_attachment(
    session_id: str,
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await verify_session_owner(db, session_id, user["id"])
    return await save_attachment(user["id"], session_id, file)


@app.delete("/api/sessions/{session_id}/attachments/{attachment_id}", status_code=204)
async def api_delete_attachment(
    session_id: str,
    attachment_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await verify_session_owner(db, session_id, user["id"])
    if not delete_attachment(user["id"], session_id, attachment_id):
        raise HTTPException(status_code=404, detail="Attachment not found.")


@app.get("/api/sessions/{session_id}/attachments/{attachment_id}/file")
async def api_get_attachment_file(
    session_id: str,
    attachment_id: str,
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    await verify_session_owner(db, session_id, user["id"])
    path, mime = get_attachment_file(user["id"], session_id, attachment_id)
    return FileResponse(path, media_type=mime, filename=path.name.split("_", 1)[-1])


# ─── Routes: Chat (SSE streaming) ────────────────────────────────────────────

@app.post("/api/chat/stream")
async def api_chat_stream(
    body: ChatRequest,
    user: dict = Depends(check_daily_message_limit),
    db: AsyncSession = Depends(get_db),
):
    """
    SSE endpoint. Opens MCP connections, runs the LangGraph agent,
    and streams tokens back to the client. Enforces daily message limit.
    """
    await verify_session_owner(db, body.session_id, user["id"])
    history = await get_session_messages(db, body.session_id)
    current_user_id = str(user["id"])

    manifest = list_attachments(current_user_id, body.session_id)
    att_meta = manifest
    if body.attachment_ids:
        allowed_ids = set(body.attachment_ids)
        att_meta = [a for a in manifest if a["id"] in allowed_ids]

    stored_user = format_stored_user_message(body.message, att_meta)
    await save_message(db, body.session_id, "user", stored_user)

    if not history:
        title = body.message[:80] + ("..." if len(body.message) > 80 else "")
        await update_session_title(db, body.session_id, title)

    logging.info("Chat stream: user_id=%s type=%s", current_user_id, type(current_user_id).__name__)
    connected_mcps = await get_connected_mcps(db, current_user_id)
    if body.mcp_ids is not None:
        allowed = set(body.mcp_ids)
        connected_mcps = [m for m in connected_mcps if m["id"] in allowed]
    session_id = body.session_id
    user_message = body.message
    attachment_context = build_agent_attachment_context(
        current_user_id,
        session_id,
        body.attachment_ids,
        history=history,
        user_message=user_message,
    )

    async def event_generator():
        full_response_parts = []

        async with get_mcp_tools(connected_mcps, user_id=current_user_id) as tools:
            router_llm = build_llm(vision=False)
            turn_intent = await classify_turn_intent(
                router_llm,
                user_message,
                attachment_ids=body.attachment_ids,
                attachment_context=attachment_context,
                tools=tools,
            )
            agent_tools, agent_attachment_context = apply_turn_routing(
                turn_intent,
                tools,
                attachment_context,
                body.attachment_ids,
            )
            async for sse_chunk in stream_agent_response(
                user_message,
                history,
                agent_tools,
                agent_attachment_context,
                attachment_ids=body.attachment_ids,
                turn_intent=turn_intent,
            ):
                yield sse_chunk

                if sse_chunk.startswith("data: "):
                    import json
                    try:
                        data = json.loads(sse_chunk[6:])
                        evt_type = data.get("type")
                        if evt_type == "token":
                            full_response_parts.append(data.get("content", ""))
                        elif evt_type == "tool_use":
                            line = f'Tool: {data["tool"]}({json.dumps(data.get("input", {}))})\n'
                            full_response_parts.append(line)
                        elif evt_type == "tool_result":
                            import base64
                            raw_content = data.get("content", "")
                            encoded = base64.b64encode(raw_content.encode("utf-8")).decode("ascii")
                            line = f'Result: {data["tool"]} -> @@JSON@@{encoded}@@END@@\n'
                            full_response_parts.append(line)
                        elif evt_type == "phase" and data.get("status") == "done":
                            detail = (data.get("detail") or "").replace("\n", " ")
                            line = (
                                f'Phase: {data.get("role", "")}|'
                                f'done|{detail}\n'
                            )
                            full_response_parts.append(line)
                    except Exception:
                        pass

        full_text = "".join(full_response_parts)
        if full_text.strip():
            async with AsyncSessionLocal() as new_db:
                await save_message(new_db, session_id, "assistant", full_text)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
        },
    )


# ─── Entry point ─────────────────────────────────────────────────────────────

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True, log_level="info")
