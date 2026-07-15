"""
chat_attachments.py - Session-scoped file/image uploads for chat context.
"""
import base64
import io
import json
import logging
import mimetypes
import re
import uuid
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import ChatSession
from app.services.document_rag import (
    chunk_text,
    select_context_text,
    save_chunks_index,
)

logger = logging.getLogger(__name__)

UPLOAD_ROOT = Path(__file__).resolve().parent / "uploads"
MAX_FILE_MB = 25
MAX_FILE_BYTES = MAX_FILE_MB * 1024 * 1024
MAX_FILES_PER_SESSION = 8
MAX_PDF_PAGES = 150
MAX_PDF_VISION_PAGES = 6
PDF_VISION_MAX_SIDE = 1280
MIN_PDF_TEXT_CHARS = 80
# Groq: base64 images in one request must stay under ~4MB total
GROQ_VISION_IMAGE_BUDGET = 3_200_000
MAX_VISION_IMAGE_BYTES = 750_000
# Per-turn character budget for document text (Groq ~128k context; leave room for chat history)
TURN_DOCUMENT_CHAR_BUDGET = 96_000
PER_FILE_CHAR_BUDGET = 58_000
MAX_IMAGES_IN_CONTEXT = 4

IMAGE_MIMES = frozenset({
    "image/jpeg", "image/png", "image/gif", "image/webp",
})
TEXT_MIMES = frozenset({
    "text/plain", "text/csv", "text/markdown", "text/html",
    "application/json", "application/xml",
})
DOC_MIMES = frozenset({
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
})
TEXT_EXTENSIONS = frozenset({
    ".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".css", ".yaml", ".yml",
    ".java", ".c", ".cpp", ".h", ".go", ".rs", ".sql", ".sh", ".bat",
    ".log", ".ini", ".env.example",
})
DOC_EXTENSIONS = frozenset({".pdf", ".docx", ".doc", ".xlsx"})

ATTACH_MARKER_RE = re.compile(r"^<!--tc-attachments:(.*?)-->\n?", re.DOTALL)

CURRENT_FILES_HEADER = (
    "## Files attached to THIS message\n"
    "Answer ONLY from the file content below for the user's current request. "
    "Do NOT use questions or answers from other documents uploaded earlier in this chat. "
    "If the user says 'solve this' or 'this quiz', they mean the file(s) below.\n"
)

FOLLOWUP_FILES_HEADER = (
    "## Files from the user's previous message in this chat\n"
    "Continue using ONLY the file content below unless the user clearly asks about something else.\n"
)


def resolve_active_attachment_ids(
    attachment_ids: list[str] | None,
    history: list[dict] | None,
) -> list[str] | None:
    """
    Decide which session attachments belong in the model context for this turn.
    - New uploads on this message → only those ids.
    - No new uploads → ids from the most recent earlier user message that had files.
    - Otherwise → no file context (avoids mixing every file ever uploaded in the session).
    """
    if attachment_ids:
        return list(attachment_ids)
    if not history:
        return None
    for msg in reversed(history):
        if msg.get("role") != "user":
            continue
        atts, _ = parse_stored_user_message(msg.get("content", ""))
        if atts:
            return [a["id"] for a in atts if a.get("id")]
    return None


def _session_dir(user_id: str, session_id: str) -> Path:
    return UPLOAD_ROOT / user_id / session_id


def _manifest_path(user_id: str, session_id: str) -> Path:
    return _session_dir(user_id, session_id) / "manifest.json"


def _load_manifest(user_id: str, session_id: str) -> list[dict]:
    path = _manifest_path(user_id, session_id)
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else []
    except Exception:
        logger.warning("Corrupt attachment manifest for session %s", session_id)
        return []


def _save_manifest(user_id: str, session_id: str, items: list[dict]) -> None:
    d = _session_dir(user_id, session_id)
    d.mkdir(parents=True, exist_ok=True)
    _manifest_path(user_id, session_id).write_text(
        json.dumps(items, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _update_manifest_entry(
    user_id: str, session_id: str, att_id: str, **updates
) -> None:
    manifest = _load_manifest(user_id, session_id)
    changed = False
    for item in manifest:
        if item["id"] == att_id:
            item.update(updates)
            changed = True
            break
    if changed:
        _save_manifest(user_id, session_id, manifest)


async def verify_session_owner(
    db: AsyncSession, session_id: str, user_id: str
) -> ChatSession:
    result = await db.execute(
        select(ChatSession).where(
            ChatSession.id == session_id,
            ChatSession.user_id == user_id,
        )
    )
    session = result.scalar_one_or_none()
    if not session:
        raise HTTPException(status_code=404, detail="Session not found.")
    return session


def _guess_mime(filename: str, declared: str | None) -> str:
    if declared and declared != "application/octet-stream":
        return declared
    guessed, _ = mimetypes.guess_type(filename)
    return guessed or "application/octet-stream"


def _is_allowed(filename: str, mime: str) -> bool:
    ext = Path(filename).suffix.lower()
    if mime in IMAGE_MIMES or mime in TEXT_MIMES or mime in DOC_MIMES:
        return True
    if ext in TEXT_EXTENSIONS or ext in DOC_EXTENSIONS:
        return True
    return False


def _pdf_page_count(data: bytes) -> int:
    try:
        import fitz
        doc = fitz.open(stream=data, filetype="pdf")
        n = min(doc.page_count, MAX_PDF_PAGES)
        doc.close()
        return n
    except Exception:
        try:
            from pypdf import PdfReader
            return min(len(PdfReader(io.BytesIO(data)).pages), MAX_PDF_PAGES)
        except Exception:
            return 0


def _extract_pdf_text_pypdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        t = page.extract_text() or ""
        if t.strip():
            parts.append(t)
    return "\n\n".join(parts)


def _extract_pdf_text_pymupdf(data: bytes) -> str:
    import fitz
    doc = fitz.open(stream=data, filetype="pdf")
    parts = []
    for i in range(min(doc.page_count, MAX_PDF_PAGES)):
        t = doc.load_page(i).get_text("text") or ""
        if t.strip():
            parts.append(t.strip())
    doc.close()
    return "\n\n".join(parts)


def _extract_pdf_text(data: bytes) -> str:
    text = _extract_pdf_text_pypdf(data)
    if len(text.strip()) < MIN_PDF_TEXT_CHARS:
        alt = _extract_pdf_text_pymupdf(data)
        if len(alt.strip()) > len(text.strip()):
            text = alt
    return text


def _pdf_needs_vision_fallback(data: bytes, extracted_text: str) -> bool:
    """Scanned or table-heavy PDFs often have little/no extractable text."""
    stripped = (extracted_text or "").strip()
    pages = _pdf_page_count(data)
    if pages == 0:
        return False
    if len(stripped) < MIN_PDF_TEXT_CHARS:
        return True
    # Very little text per page → likely image-based pages
    return len(stripped) < pages * 50


def _to_vision_data_url(raw_bytes: bytes) -> str | None:
    """Compress image for Groq vision limits (~4MB total request, per-image caps)."""
    try:
        from PIL import Image

        img = Image.open(io.BytesIO(raw_bytes))
        if img.mode in ("RGBA", "LA", "P"):
            background = Image.new("RGB", img.size, (255, 255, 255))
            if img.mode == "P":
                img = img.convert("RGBA")
            if img.mode in ("RGBA", "LA"):
                background.paste(img, mask=img.split()[-1])
            else:
                background.paste(img)
            img = background
        elif img.mode != "RGB":
            img = img.convert("RGB")

        img.thumbnail((PDF_VISION_MAX_SIDE, PDF_VISION_MAX_SIDE))

        jpeg = raw_bytes
        for quality in (78, 68, 58, 48, 40):
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=quality, optimize=True)
            candidate = buf.getvalue()
            if len(candidate) <= MAX_VISION_IMAGE_BYTES:
                jpeg = candidate
                break
            jpeg = candidate

        b64 = base64.b64encode(jpeg).decode("ascii")
        return f"data:image/jpeg;base64,{b64}"
    except Exception as exc:
        logger.warning("Vision image compress failed: %s", exc)
        return None


def _render_pdf_pages_as_data_urls(data: bytes, max_pages: int = MAX_PDF_VISION_PAGES) -> list[str]:
    """Render PDF pages as compressed JPEG data URLs (fits Groq 4MB vision limit)."""
    import fitz

    doc = fitz.open(stream=data, filetype="pdf")
    urls: list[str] = []
    budget = GROQ_VISION_IMAGE_BUDGET
    limit = min(doc.page_count, max_pages, MAX_IMAGES_IN_CONTEXT)

    for i in range(limit):
        if budget < 50_000:
            break
        page = doc.load_page(i)
        rect = page.rect
        if rect.width <= 0 or rect.height <= 0:
            continue
        zoom = min(PDF_VISION_MAX_SIDE / rect.width, PDF_VISION_MAX_SIDE / rect.height, 2.0)
        mat = fitz.Matrix(zoom, zoom)
        pix = page.get_pixmap(matrix=mat, alpha=False)
        url = _to_vision_data_url(pix.tobytes("png"))
        if not url or len(url) > budget:
            continue
        urls.append(url)
        budget -= len(url)

    doc.close()
    return urls


def _extract_xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"## Sheet: {sheet.title}")
        for row in sheet.iter_rows(max_row=500, values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n\n".join(parts)


def _extract_docx_text(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _extract_text(data: bytes, mime: str, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    try:
        if mime == "application/pdf" or ext == ".pdf":
            return _extract_pdf_text(data)  # caller checks vision fallback separately for PDFs
        if (
            mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or ext == ".docx"
        ):
            return _extract_docx_text(data)
        if (
            mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            or ext == ".xlsx"
        ):
            return _extract_xlsx_text(data)
    except Exception as exc:
        logger.warning("Document extract failed for %s: %s", filename, exc)
        return ""

    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return ""


def _text_cache_path(user_id: str, session_id: str, att_id: str) -> Path:
    return _session_dir(user_id, session_id) / f"{att_id}_extracted.txt"


def _chunks_index_path(user_id: str, session_id: str, att_id: str) -> Path:
    return _session_dir(user_id, session_id) / f"{att_id}_chunks.json"


def _persist_full_text(
    user_id: str, session_id: str, att_id: str, raw: str
) -> tuple[int, int]:
    """Save full extracted text + chunk index on disk (Claude-style full ingest)."""
    raw = (raw or "").strip()
    cache = _text_cache_path(user_id, session_id, att_id)
    if raw:
        cache.write_text(raw, encoding="utf-8")
        chunks = chunk_text(raw)
        save_chunks_index(_chunks_index_path(user_id, session_id, att_id), chunks)
        return len(raw), len(chunks)
    return 0, 0


def _load_full_text(user_id: str, session_id: str, entry: dict) -> str:
    att_id = entry["id"]
    cache = _text_cache_path(user_id, session_id, att_id)
    if cache.exists():
        return cache.read_text(encoding="utf-8")

    if entry.get("kind") == "image":
        return ""

    try:
        data = _read_attachment_bytes(user_id, session_id, att_id)
        raw = _extract_text(data, entry.get("mime", ""), entry.get("name", "file"))
        if raw.strip():
            char_count, chunk_count = _persist_full_text(
                user_id, session_id, att_id, raw
            )
            _update_manifest_entry(
                user_id,
                session_id,
                att_id,
                char_count=char_count,
                chunk_count=chunk_count,
            )
            entry["char_count"] = char_count
            entry["chunk_count"] = chunk_count
        return raw
    except Exception as exc:
        logger.warning("Load full text failed for %s: %s", entry.get("name"), exc)
        return entry.get("text_excerpt") or ""


def _kind_for(mime: str, filename: str) -> str:
    if mime in IMAGE_MIMES:
        return "image"
    return "document"


def _read_attachment_bytes(user_id: str, session_id: str, att_id: str) -> bytes:
    path, _ = get_attachment_file(user_id, session_id, att_id)
    return path.read_bytes()


def resolve_document_context(
    user_id: str,
    session_id: str,
    entry: dict,
    user_query: str,
    char_budget: int,
) -> tuple[str, dict]:
    """Select full doc or relevant excerpts for the user's question (RAG-style)."""
    full_text = _load_full_text(user_id, session_id, entry)
    if not full_text.strip():
        return "", {"mode": "empty"}
    return select_context_text(
        full_text,
        user_query,
        char_budget,
        filename=entry.get("name", "file"),
    )


async def save_attachment(
    user_id: str,
    session_id: str,
    upload: UploadFile,
) -> dict:
    manifest = _load_manifest(user_id, session_id)
    if len(manifest) >= MAX_FILES_PER_SESSION:
        raise HTTPException(
            status_code=400,
            detail=f"Maximum {MAX_FILES_PER_SESSION} files per chat.",
        )

    filename = (upload.filename or "file").replace("\\", "/").split("/")[-1]
    if not filename or filename in (".", ".."):
        raise HTTPException(status_code=400, detail="Invalid file name.")

    data = await upload.read()
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(
            status_code=400,
            detail=f"File too large (max {MAX_FILE_MB} MB).",
        )
    if not data:
        raise HTTPException(status_code=400, detail="Empty file.")

    mime = _guess_mime(filename, upload.content_type)
    if not _is_allowed(filename, mime):
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Use images, PDF, Word, Excel, or text/code files.",
        )

    att_id = str(uuid.uuid4())
    safe_name = f"{att_id}_{filename}"
    dest_dir = _session_dir(user_id, session_id)
    dest_dir.mkdir(parents=True, exist_ok=True)
    file_path = dest_dir / safe_name
    file_path.write_bytes(data)

    kind = _kind_for(mime, filename)
    char_count = 0
    chunk_count = 0
    pdf_vision = False
    if kind == "document":
        raw = _extract_text(data, mime, filename)
        if mime == "application/pdf" or filename.lower().endswith(".pdf"):
            pdf_vision = _pdf_needs_vision_fallback(data, raw)
        if raw.strip() and not pdf_vision:
            char_count, chunk_count = _persist_full_text(
                user_id, session_id, att_id, raw
            )
        elif pdf_vision:
            logger.info(
                "PDF %s: using vision fallback (%d pages, %d chars text)",
                filename,
                _pdf_page_count(data),
                len(raw.strip()),
            )
        elif not raw.strip():
            logger.warning(
                "No text extracted from %s (mime=%s).",
                filename,
                mime,
            )

    entry = {
        "id": att_id,
        "name": filename,
        "mime": mime,
        "kind": kind,
        "size": len(data),
        "char_count": char_count,
        "chunk_count": chunk_count,
        "pdf_vision": pdf_vision,
    }
    manifest.append(entry)
    _save_manifest(user_id, session_id, manifest)
    return {
        "id": att_id,
        "name": filename,
        "mime": mime,
        "kind": kind,
        "size": len(data),
        "has_text": bool(char_count),
        "char_count": char_count,
        "pdf_vision": pdf_vision,
    }


def list_attachments(user_id: str, session_id: str) -> list[dict]:
    return [
        {
            "id": a["id"],
            "name": a["name"],
            "mime": a["mime"],
            "kind": a["kind"],
            "size": a["size"],
        }
        for a in _load_manifest(user_id, session_id)
    ]


def delete_attachment(user_id: str, session_id: str, attachment_id: str) -> bool:
    manifest = _load_manifest(user_id, session_id)
    found = None
    remaining = []
    for item in manifest:
        if item["id"] == attachment_id:
            found = item
        else:
            remaining.append(item)
    if not found:
        return False
    _save_manifest(user_id, session_id, remaining)
    d = _session_dir(user_id, session_id)
    for p in d.glob(f"{attachment_id}_*"):
        try:
            p.unlink()
        except OSError:
            pass
    for suffix in ("_extracted.txt", "_chunks.json"):
        extra = d / f"{attachment_id}{suffix}"
        if extra.exists():
            try:
                extra.unlink()
            except OSError:
                pass
    return True


def get_attachment_file(user_id: str, session_id: str, attachment_id: str) -> tuple[Path, str]:
    manifest = _load_manifest(user_id, session_id)
    entry = next((a for a in manifest if a["id"] == attachment_id), None)
    if not entry:
        raise HTTPException(status_code=404, detail="Attachment not found.")
    d = _session_dir(user_id, session_id)
    matches = list(d.glob(f"{attachment_id}_*"))
    if not matches:
        raise HTTPException(status_code=404, detail="Attachment file missing.")
    return matches[0], entry["mime"]


def clear_session_attachments(user_id: str, session_id: str) -> None:
    d = _session_dir(user_id, session_id)
    if d.exists():
        import shutil
        shutil.rmtree(d, ignore_errors=True)


def format_stored_user_message(text: str, attachments: list[dict]) -> str:
    if not attachments:
        return text
    meta = json.dumps(
        [{"id": a["id"], "name": a["name"], "kind": a.get("kind", "text")} for a in attachments],
        ensure_ascii=False,
    )
    return f"<!--tc-attachments:{meta}-->\n{text}"


def parse_stored_user_message(content: str) -> tuple[list[dict], str]:
    m = ATTACH_MARKER_RE.match(content)
    if not m:
        return [], content
    try:
        attachments = json.loads(m.group(1))
        if not isinstance(attachments, list):
            attachments = []
    except json.JSONDecodeError:
        attachments = []
    body = content[m.end():]
    return attachments, body


def build_agent_attachment_context(
    user_id: str,
    session_id: str,
    attachment_ids: list[str] | None = None,
    history: list[dict] | None = None,
    user_message: str = "",
) -> dict[str, Any]:
    """
    Build text prefix and image data URLs for the active turn only (not every file in session).
    Re-extracts document text from disk when cached excerpt is empty.
    """
    empty = {
        "text_prefix": "",
        "images": [],
        "use_vision": False,
        "attachments": [],
        "has_files": False,
        "is_current_turn": bool(attachment_ids),
    }
    manifest = _load_manifest(user_id, session_id)
    if not manifest:
        return empty

    active_ids = resolve_active_attachment_ids(attachment_ids, history)
    if not active_ids:
        return empty

    active_set = set(active_ids)
    entries = [e for e in manifest if e["id"] in active_set]
    if not entries:
        return empty

    is_current_turn = bool(attachment_ids)
    header = CURRENT_FILES_HEADER if is_current_turn else FOLLOWUP_FILES_HEADER
    text_parts: list[str] = [header]
    total_chars = len(header)
    images: list[str] = []
    image_count = 0
    has_file_content = False

    doc_entries = [e for e in entries if e.get("kind") == "document"]
    per_file_budget = min(
        PER_FILE_CHAR_BUDGET,
        TURN_DOCUMENT_CHAR_BUDGET // max(len(doc_entries), 1),
    )

    for entry in entries:
        name = entry.get("name", "file")
        kind = entry.get("kind", "document")

        if kind == "document":
            is_pdf = (
                entry.get("mime") == "application/pdf"
                or name.lower().endswith(".pdf")
            )
            use_pdf_vision = bool(entry.get("pdf_vision"))
            context_body = ""
            rag_meta: dict = {}

            if not use_pdf_vision:
                remaining = TURN_DOCUMENT_CHAR_BUDGET - total_chars
                budget = min(per_file_budget, remaining)
                if budget < 2000:
                    budget = max(remaining, 2000)
                context_body, rag_meta = resolve_document_context(
                    user_id, session_id, entry, user_message, budget
                )

            # Scanned PDF or empty extract → send page images to vision model
            if is_pdf and (use_pdf_vision or not context_body.strip()):
                try:
                    pdf_bytes = _read_attachment_bytes(
                        user_id, session_id, entry["id"]
                    )
                    if not entry.get("pdf_vision") and _pdf_needs_vision_fallback(
                        pdf_bytes, context_body
                    ):
                        use_pdf_vision = True
                    page_urls = _render_pdf_pages_as_data_urls(pdf_bytes)
                    for url in page_urls:
                        if image_count >= MAX_IMAGES_IN_CONTEXT:
                            break
                        images.append(url)
                        image_count += 1
                    if page_urls:
                        has_file_content = True
                        note = (
                            f"### File: {name} (PDF — {len(page_urls)} page(s) as images)\n"
                            "This PDF is scanned or image-based. Read the page images below "
                            "and explain the schedule/content in detail.\n"
                        )
                        if context_body.strip():
                            note += f"\nPartial extracted text:\n```\n{context_body[:8000]}\n```\n"
                        text_parts.append(note)
                        total_chars += len(note)
                except Exception as exc:
                    logger.warning("PDF vision render failed for %s: %s", name, exc)

            if context_body.strip() and not use_pdf_vision:
                mode = rag_meta.get("mode", "full")
                label = "complete document" if mode == "full" else "relevant excerpts"
                block = f"### File: {name} ({label})\n```\n{context_body}\n```"
                has_file_content = True
                if total_chars + len(block) > TURN_DOCUMENT_CHAR_BUDGET:
                    block = block[: max(0, TURN_DOCUMENT_CHAR_BUDGET - total_chars)] + "\n…"
                text_parts.append(block)
                total_chars += len(block)
            elif not has_file_content and not images:
                block = (
                    f"### File: {name}\n"
                    f"(Attached {entry.get('mime', 'file')}; text could not be extracted. "
                    "Tell the user the file is attached but readable text was not available.)"
                )
                text_parts.append(block)
                total_chars += len(block)

        elif kind == "image":
            if image_count < MAX_IMAGES_IN_CONTEXT:
                try:
                    path, mime = get_attachment_file(user_id, session_id, entry["id"])
                    url = _to_vision_data_url(path.read_bytes())
                    if url:
                        images.append(url)
                    image_count += 1
                    text_parts.append(f"### Image: {name}\n(see image below)")
                    has_file_content = True
                except HTTPException:
                    text_parts.append(f"### Image: {name}\n(image file attached)")

    text_prefix = "\n\n".join(text_parts) if has_file_content or len(text_parts) > 1 else ""
    return {
        "text_prefix": text_prefix,
        "images": images,
        "use_vision": bool(images),
        "attachments": entries,
        "has_files": bool(entries),
        "is_current_turn": is_current_turn,
    }
