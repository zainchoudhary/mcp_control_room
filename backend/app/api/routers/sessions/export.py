"""Export chat sessions to Word document."""
import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.auth_utils import get_current_user
from app.db.database import get_all_user_sessions_with_messages
from app.db.config import get_db

router = APIRouter(prefix="/api/sessions", tags=["Sessions"])


@router.get("/export")
async def export_sessions(
    user: dict = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Export all chat sessions as a Word .docx file."""
    sessions = await get_all_user_sessions_with_messages(db, user["id"])

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title_para = doc.add_heading("Chat Export \u2014 ToolChain AI", level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    username = user.get("username") or user.get("email") or "User"
    doc.add_paragraph(f"Exported by: {username}")
    doc.add_paragraph(f"Date: {datetime.utcnow().strftime('%B %d, %Y at %H:%M UTC')}")
    doc.add_paragraph(f"Total sessions: {len(sessions)}")

    if not sessions:
        doc.add_paragraph("\nNo chat sessions found.")
    else:
        for idx, session in enumerate(sessions, 1):
            created = session.get("created_at", "")
            if created:
                try:
                    dt = datetime.fromisoformat(str(created).replace("Z", "+00:00"))
                    date_str = dt.strftime("%B %d, %Y")
                except Exception:
                    date_str = str(created)[:10]
            else:
                date_str = "Unknown date"

            title_text = session.get("title") or f"Session {idx}"
            doc.add_heading(f"{title_text} ({date_str})", level=2)

            messages = session.get("messages", [])
            if not messages:
                doc.add_paragraph("(No messages)")
            else:
                for msg in messages:
                    role = msg.get("role", "unknown")
                    content = msg.get("content", "")
                    label = "You" if role == "user" else "Assistant" if role == "assistant" else role.capitalize()

                    p = doc.add_paragraph()
                    run = p.add_run(f"[{label}]: ")
                    run.bold = True
                    if role == "user":
                        run.font.color.rgb = RGBColor(0x33, 0x33, 0xCC)
                    else:
                        run.font.color.rgb = RGBColor(0x10, 0xA3, 0x7F)
                    p.add_run(content)

            if idx < len(sessions):
                doc.add_paragraph("\u2500" * 50)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=toolchain_chats_export.docx"},
    )
